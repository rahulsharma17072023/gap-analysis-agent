"""
Mordor Intelligence - Competitor Title Gap Analysis Agent
Simple version without external dependencies issues
"""

import os
import re
from datetime import datetime
from io import BytesIO

from flask import Flask, request, jsonify, send_file
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side

app = Flask(__name__)

# Simple CORS support
@app.after_request
def after_request(response):
    response.headers.add('Access-Control-Allow-Origin', '*')
    response.headers.add('Access-Control-Allow-Headers', 'Content-Type')
    response.headers.add('Access-Control-Allow-Methods', 'GET, POST, OPTIONS')
    return response

# ==================== CONFIGURATION ====================

COMPETITOR_DATA = {
    "Allied Market Research": ["Global Healthcare Market Size & Forecast 2030", "U.S. Retail Fashion Market Outlook 2026", "Artificial Intelligence in Manufacturing Market Analysis", "Cloud Computing Market 2025-2032", "5G Technology Market Growth Projections"],
    "Global Data Inc.": ["Healthcare Diagnostics Market 2026", "Fintech Innovation Market Size & Forecast", "Supply Chain Logistics Market 2025-2030", "Renewable Energy Market Analysis", "Smart City Technology Market"],
    "Technavio": ["Market Research on Global Healthcare IT", "Retail E-commerce Market Forecast", "Industrial Automation Market Growth", "Consumer Electronics Market Analysis", "Food and Beverage Technology Market"],
    "IMARC": ["Market Report: Pharmaceutical Industry 2026", "Global Retail Market Size and Trends", "Agriculture Technology Market Analysis", "Construction Technology Market 2025", "Healthcare Services Market Forecast"],
    "Armstrong & Associates": ["Logistics and Supply Chain Market", "Warehouse Management Systems Market", "Transportation Technology Market Analysis", "Freight and Shipping Market 2026", "Last-Mile Delivery Market Report"],
    "Future Market Insights": ["Luxury Market Analysis 2026", "Premium Food Market Report", "High-End Fashion Market Forecast", "Luxury Automotive Market", "Fine Dining Market Analysis"],
    "Grand View Research": ["Market Size, Share & Trends Analysis: Healthcare", "Global Pharmaceutical Market Forecast", "Medical Devices Market Report 2026", "Biotech Market Analysis and Outlook", "Clinical Diagnostics Market"],
    "Kentley Insights": ["Consumer Goods Market Trends 2026", "Household Products Market Analysis", "Personal Care Market Forecast", "Food Ingredients Market Report", "Beverage Industry Market 2026"],
    "The Insight Partners": ["BFSI Market Analysis 2026", "Banking and Finance Technology Market", "Insurance Technology Market Forecast", "Investment Management Solutions Market", "Risk Management Software Market"],
    "Markets and Markets": ["Semiconductor Market Size and Forecast", "Consumer Electronics Market 2026", "Display Technology Market Report", "Memory Devices Market Analysis", "Microprocessor Market Forecast"],
}

STOPWORDS = {"market", "size", "share", "forecast", "analysis", "trends", "report", "outlook", "demand", "a", "the", "and", "or", "in", "of", "by", "for", "to", "from", "with", "on", "at", "is", "are", "was", "were", "been", "be", "have", "has", "do", "does", "did", "will", "would"}

CLUSTER_KEYWORDS = {
    "HPI&M": {"health", "healthcare", "medical", "pharma", "pharmaceutical", "drug", "medicine", "hospital", "clinic", "patient", "disease", "treatment", "therapy", "insurance", "diagnostic"},
    "RCM1": {"retail", "consumer", "fashion", "apparel", "clothing", "beauty", "cosmetics", "personal care", "food", "beverage", "grocery", "shopping", "brand", "product", "goods"},
    "Logistics": {"logistics", "supply", "chain", "transport", "shipping", "freight", "delivery", "warehouse", "distribution", "inventory", "port", "cargo", "fleet", "vehicle"},
    "FSII&PS": {"financial", "banking", "semiconductor", "software", "technology", "tech", "it", "digital", "computer", "network", "data", "payment", "fintech", "cryptocurrency", "blockchain"}
}

SIMILARITY_THRESHOLD = 0.40

# ==================== HELPER CLASSES ====================

class TitleNormalizer:
    @staticmethod
    def normalize(title):
        if not title or not isinstance(title, str):
            return set()
        title = title.lower()
        title = re.sub(r'\b(20|21)\d{2}(-\d{2})?\b', '', title)
        title = re.sub(r'[^a-z0-9\s]', ' ', title)
        tokens = set(title.split())
        tokens = tokens - STOPWORDS
        tokens.discard('')
        return tokens

class TitleClassifier:
    @staticmethod
    def classify(title):
        if not title:
            return "Unclassified"
        title_lower = title.lower()
        scores = {cluster: sum(1 for kw in keywords if kw in title_lower) for cluster, keywords in CLUSTER_KEYWORDS.items()}
        return max(scores, key=scores.get) if max(scores.values()) > 0 else "Unclassified"

class TitleMatcher:
    @staticmethod
    def jaccard_similarity(tokens1, tokens2):
        if not tokens1 or not tokens2:
            return 0.0
        intersection = len(tokens1 & tokens2)
        union = len(tokens1 | tokens2)
        return intersection / union if union > 0 else 0.0

    @staticmethod
    def match(competitor_title, mordor_titles, threshold=SIMILARITY_THRESHOLD):
        comp_tokens = TitleNormalizer.normalize(competitor_title)
        best_match = None
        best_score = 0.0
        for mordor_title in mordor_titles:
            mordor_tokens = TitleNormalizer.normalize(mordor_title)
            score = TitleMatcher.jaccard_similarity(comp_tokens, mordor_tokens)
            if score > best_score:
                best_score = score
                best_match = mordor_title
        return (best_match, best_score) if best_score >= threshold else (None, best_score)

# ==================== ANALYSIS ENGINE ====================

class GapAnalysisEngine:
    @staticmethod
    def analyze(mordor_titles, competitors_list):
        active_titles = [t for t in mordor_titles if t.get('Status') == 'Active']
        if not active_titles:
            raise ValueError("No active titles found. Make sure Status column has 'Active' values.")

        mordor_by_cluster = {}
        for title_obj in active_titles:
            cluster = title_obj.get('Practice_Cluster') or TitleClassifier.classify(title_obj.get('Title', ''))
            if cluster not in mordor_by_cluster:
                mordor_by_cluster[cluster] = []
            mordor_by_cluster[cluster].append(title_obj.get('Title', ''))

        gaps, covered, competitor_stats = [], [], {}

        for comp_name in competitors_list:
            if comp_name not in COMPETITOR_DATA:
                continue
            comp_titles = COMPETITOR_DATA[comp_name]
            comp_gap_count = comp_covered_count = 0

            for comp_title in comp_titles:
                comp_cluster = TitleClassifier.classify(comp_title)
                cluster_titles = mordor_by_cluster.get(comp_cluster, [])

                if cluster_titles:
                    matched, score = TitleMatcher.match(comp_title, cluster_titles, SIMILARITY_THRESHOLD)
                    if matched:
                        covered.append({'competitor': comp_name, 'competitor_title': comp_title, 'mordor_title': matched, 'cluster': comp_cluster, 'similarity_score': round(score, 3), 'confidence': 'High' if score >= 0.60 else 'Medium' if score >= 0.40 else 'Low'})
                        comp_covered_count += 1
                    else:
                        gaps.append({'competitor': comp_name, 'competitor_title': comp_title, 'cluster': comp_cluster, 'similarity_score': round(score, 3), 'confidence': 'High' if score >= 0.60 else 'Medium' if score >= 0.40 else 'Low'})
                        comp_gap_count += 1
                else:
                    gaps.append({'competitor': comp_name, 'competitor_title': comp_title, 'cluster': comp_cluster, 'similarity_score': 0.0, 'confidence': 'Low'})
                    comp_gap_count += 1

            competitor_stats[comp_name] = {'total_titles': len(comp_titles), 'gaps': comp_gap_count, 'covered': comp_covered_count}

        total_comp_titles = sum(len(COMPETITOR_DATA[c]) for c in competitors_list if c in COMPETITOR_DATA)
        total_gaps = len(gaps)
        total_covered = len(covered)
        gap_pct = round((total_gaps / total_comp_titles * 100) if total_comp_titles > 0 else 0, 1)

        return {'gaps': gaps, 'covered': covered, 'summary': {'mordor_titles_active': len(active_titles), 'competitors_analyzed': len(competitors_list), 'total_competitor_titles': total_comp_titles, 'gaps_count': total_gaps, 'covered_count': total_covered, 'gap_percentage': gap_pct, 'clusters': list(mordor_by_cluster.keys())}, 'competitor_stats': competitor_stats}

# ==================== EXCEL REPORT ====================

def generate_excel_report(analysis, date_start, date_end):
    wb = Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill(start_color="667eea", end_color="667eea", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=12)
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    # SUMMARY
    ws = wb.create_sheet("Summary")
    summary = analysis['summary']
    row = 1
    ws[f'A{row}'] = "Competitor Title Gap Analysis Report"
    ws[f'A{row}'].font = title_font
    row += 2
    ws[f'A{row}'] = "Date Range:"
    ws[f'B{row}'] = f"{date_start} to {date_end}"
    row += 1
    ws[f'A{row}'] = "Analysis Date:"
    ws[f'B{row}'] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    row += 2
    ws[f'A{row}'] = "Key Metrics:"
    ws[f'A{row}'].font = title_font
    row += 1

    for metric, value in [("Mordor Titles (Active)", summary['mordor_titles_active']), ("Competitors Analyzed", summary['competitors_analyzed']), ("Total Competitor Titles", summary['total_competitor_titles']), ("Covered by Mordor", summary['covered_count']), ("Gap / Whitespace", summary['gaps_count']), ("Gap Percentage", f"{summary['gap_percentage']}%")]:
        ws[f'A{row}'] = metric
        ws[f'B{row}'] = value
        row += 1

    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 20

    # GAP DETAIL
    ws_gap = wb.create_sheet("Gap Detail")
    headers = ["Competitor", "Gap Title", "Cluster", "Similarity Score", "Confidence"]
    for col, header in enumerate(headers, 1):
        cell = ws_gap.cell(row=1, column=col)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border

    for idx, gap in enumerate(analysis['gaps'], 2):
        ws_gap.cell(row=idx, column=1).value = gap['competitor']
        ws_gap.cell(row=idx, column=2).value = gap['competitor_title']
        ws_gap.cell(row=idx, column=3).value = gap['cluster']
        ws_gap.cell(row=idx, column=4).value = gap['similarity_score']
        ws_gap.cell(row=idx, column=5).value = gap['confidence']
        for col in range(1, 6):
            ws_gap.cell(row=idx, column=col).border = border

    for col, width in enumerate(['A', 'B', 'C', 'D', 'E'], 1):
        ws_gap.column_dimensions[chr(64+col)].width = [25, 50, 15, 18, 15][col-1]

    # COVERED
    ws_cov = wb.create_sheet("Already Covered")
    headers = ["Competitor", "Competitor Title", "Matched Mordor Title", "Cluster", "Similarity Score", "Confidence"]
    for col, header in enumerate(headers, 1):
        cell = ws_cov.cell(row=1, column=col)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border

    for idx, item in enumerate(analysis['covered'], 2):
        ws_cov.cell(row=idx, column=1).value = item['competitor']
        ws_cov.cell(row=idx, column=2).value = item['competitor_title']
        ws_cov.cell(row=idx, column=3).value = item['mordor_title']
        ws_cov.cell(row=idx, column=4).value = item['cluster']
        ws_cov.cell(row=idx, column=5).value = item['similarity_score']
        ws_cov.cell(row=idx, column=6).value = item['confidence']
        for col in range(1, 7):
            ws_cov.cell(row=idx, column=col).border = border

    for col, width in enumerate(['A', 'B', 'C', 'D', 'E', 'F'], 1):
        ws_cov.column_dimensions[chr(64+col)].width = [25, 40, 40, 15, 18, 15][col-1]

    # METADATA
    ws_meta = wb.create_sheet("Metadata")
    row = 1
    ws_meta[f'A{row}'] = "Metadata & Configuration"
    ws_meta[f'A{row}'].font = title_font
    row += 2

    for key, value in [("Run Timestamp", datetime.now().isoformat()), ("Date Range Start", date_start), ("Date Range End", date_end), ("Competitors Analyzed", summary['competitors_analyzed']), ("Similarity Threshold", SIMILARITY_THRESHOLD), ("Total Gap Titles", summary['gaps_count']), ("Total Covered Titles", summary['covered_count'])]:
        ws_meta[f'A{row}'] = key
        ws_meta[f'B{row}'] = value
        row += 1

    ws_meta.column_dimensions['A'].width = 25
    ws_meta.column_dimensions['B'].width = 40

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output

# ==================== ROUTES ====================

@app.route('/api/competitors', methods=['GET'])
def get_competitors():
    return jsonify({'competitors': list(COMPETITOR_DATA.keys()), 'count': len(COMPETITOR_DATA)})

@app.route('/api/analyze', methods=['POST'])
def analyze():
    try:
        if 'mordor_file' not in request.files:
            return jsonify({'error': 'No file uploaded'}), 400
        if 'competitors' not in request.form:
            return jsonify({'error': 'No competitors selected'}), 400
        if 'dateStart' not in request.form or 'dateEnd' not in request.form:
            return jsonify({'error': 'Date range required'}), 400

        file = request.files['mordor_file']
        competitors = request.form.getlist('competitors')
        date_start = request.form.get('dateStart')
        date_end = request.form.get('dateEnd')

        if file.filename == '':
            return jsonify({'error': 'No file selected'}), 400
        if not file.filename.endswith(('.xlsx', '.xls')):
            return jsonify({'error': 'File must be Excel (.xlsx or .xls)'}), 400
        if not competitors:
            return jsonify({'error': 'At least one competitor must be selected'}), 400

        try:
            excel_data = pd.read_excel(file)
        except Exception as e:
            return jsonify({'error': f'Failed to read Excel file: {str(e)}'}), 400

        if 'Title' not in excel_data.columns or 'Status' not in excel_data.columns:
            return jsonify({'error': 'Excel must have "Title" and "Status" columns'}), 400

        mordor_titles = excel_data.to_dict('records')
        analysis = GapAnalysisEngine.analyze(mordor_titles, competitors)
        excel_file = generate_excel_report(analysis, date_start, date_end)
        filename = f"Gap_Analysis_{date_start.replace('-', '')}.xlsx"

        return send_file(excel_file, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name=filename)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/health', methods=['GET'])
def health():
    return jsonify({'status': 'healthy', 'timestamp': datetime.now().isoformat()})

@app.route('/', methods=['GET'])
def index():
    return '''<!DOCTYPE html><html><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width, initial-scale=1.0"><title>Gap Analysis Platform</title><style>*{margin:0;padding:0;box-sizing:border-box}body{font-family:'Segoe UI',Tahoma,Geneva,Verdana,sans-serif;background:linear-gradient(135deg,#667eea 0%,#764ba2 100%);min-height:100vh;padding:20px}.container{max-width:900px;margin:0 auto;background:white;border-radius:10px;box-shadow:0 20px 60px rgba(0,0,0,.3)}.header{background:linear-gradient(135deg,#667eea 0%,#764ba2 100%);color:white;padding:40px 30px;text-align:center}.header h1{font-size:2.5em;margin-bottom:10px}.header p{font-size:1.1em;opacity:.9}.content{padding:40px 30px}.form-section{margin-bottom:30px}.form-section h2{font-size:1.3em;color:#333;margin-bottom:15px;padding-bottom:10px;border-bottom:2px solid #667eea}.form-group{margin-bottom:20px}.form-group label{display:block;font-weight:600;color:#333;margin-bottom:8px}.form-group input,.form-group select{width:100%;padding:12px;border:1px solid #ddd;border-radius:5px;font-size:.95em;font-family:inherit}.date-range{display:grid;grid-template-columns:1fr 1fr;gap:15px}.file-upload-area{border:2px dashed #667eea;border-radius:5px;padding:30px;text-align:center;cursor:pointer}.file-upload-area:hover{background:#f0f1ff;border-color:#764ba2}.button-group{display:grid;grid-template-columns:1fr 1fr;gap:15px}button{padding:12px 30px;border:none;border-radius:5px;font-size:1em;font-weight:600;cursor:pointer}.btn-submit{background:linear-gradient(135deg,#667eea 0%,#764ba2 100%);color:white}.btn-submit:disabled{opacity:.5;cursor:not-allowed}.btn-reset{background:#f0f0f0;color:#333}.message{padding:12px;border-radius:5px;margin-bottom:15px;display:none}.error{background:#fee;border:1px solid #fcc;color:#c33}.success{background:#efe;border:1px solid #cfc;color:#3c3}.progress-section{display:none;margin-top:20px;padding:20px;background:#f9f9f9;border-radius:5px}.progress-section.active{display:block}.progress-bar{width:100%;height:10px;background:#ddd;border-radius:5px;overflow:hidden}.progress-fill{height:100%;background:linear-gradient(90deg,#667eea 0%,#764ba2 100%);width:100%;animation:pulse 1s infinite}@keyframes pulse{0%,100%{opacity:1}50%{opacity:.7}}.results-section{display:none;margin-top:20px;padding:20px;background:#efe;border:1px solid #cfc;border-radius:5px;text-align:center}.results-section.active{display:block}.results-section a{display:inline-block;background:linear-gradient(135deg,#667eea 0%,#764ba2 100%);color:white;padding:12px 30px;border-radius:5px;text-decoration:none;font-weight:600;margin-top:10px}@media (max-width:600px){.date-range{grid-template-columns:1fr}.button-group{grid-template-columns:1fr}}</style></head><body><div class="container"><div class="header"><h1>📊 Gap Analysis Platform</h1><p>Competitor Title Gap Analysis for Mordor Intelligence</p></div><div class="content"><div id="errorMessage" class="message error"></div><div id="successMessage" class="message success"></div><form id="analysisForm"><div class="form-section"><h2>📅 Date Range</h2><div class="date-range"><div class="form-group"><label for="dateStart">Start Date</label><input type="date" id="dateStart" name="dateStart" required></div><div class="form-group"><label for="dateEnd">End Date</label><input type="date" id="dateEnd" name="dateEnd" required></div></div></div><div class="form-section"><h2>🏢 Competitors</h2><div class="form-group"><label for="competitors">Select Competitors</label><select id="competitors" name="competitors" multiple required style="height:150px;"></select></div></div><div class="form-section"><h2>📁 Mordor Titles Portfolio</h2><div class="form-group"><label for="mordorFile">Upload Excel File (XLSX/XLS)</label><div class="file-upload-area" id="fileUploadArea"><p>📥 Drag and drop your Excel file here or click to browse</p><input type="file" id="mordorFile" name="mordor_file" accept=".xlsx,.xls" required style="display:none;"><div id="fileName" style="margin-top:10px;color:#667eea;font-weight:600;"></div></div><p style="font-size:.85em;color:#999;margin-top:10px;">Required columns: Title, Status. Optional: Practice_Cluster</p></div></div><div class="button-group"><button type="submit" class="btn-submit" id="submitBtn">🚀 Analyze</button><button type="reset" class="btn-reset">🔄 Clear</button></div></form><div id="progressSection" class="progress-section"><h3>Analysis in Progress...</h3><div class="progress-bar"><div class="progress-fill"></div></div><p style="margin-top:10px;color:#666;">Processing your Mordor portfolio and analyzing gaps...</p></div><div id="resultsSection" class="results-section"><h3>✅ Analysis Complete!</h3><p>Your gap analysis report is ready for download.</p><a id="downloadLink" href="#" download>📥 Download Excel Report</a></div></div></div><script>fetch('/api/competitors').then(r=>r.json()).then(data=>{const select=document.getElementById('competitors');data.competitors.forEach(comp=>{const option=document.createElement('option');option.value=comp;option.textContent=comp;select.appendChild(option)})});const today=new Date();const firstDay=new Date(today.getFullYear(),today.getMonth(),1);const lastDay=new Date(today.getFullYear(),today.getMonth()+1,0);document.getElementById('dateStart').valueAsDate=firstDay;document.getElementById('dateEnd').valueAsDate=lastDay;const fileUploadArea=document.getElementById('fileUploadArea');const fileInput=document.getElementById('mordorFile');const fileName=document.getElementById('fileName');fileUploadArea.addEventListener('click',()=>fileInput.click());fileUploadArea.addEventListener('dragover',(e)=>{e.preventDefault();fileUploadArea.style.borderColor='#764ba2';fileUploadArea.style.background='#e8e9ff'});fileUploadArea.addEventListener('dragleave',()=>{fileUploadArea.style.borderColor='#667eea';fileUploadArea.style.background='#f8f9ff'});fileUploadArea.addEventListener('drop',(e)=>{e.preventDefault();fileUploadArea.style.borderColor='#667eea';fileUploadArea.style.background='#f8f9ff';fileInput.files=e.dataTransfer.files;updateFileName()});fileInput.addEventListener('change',updateFileName);function updateFileName(){if(fileInput.files.length>0){fileName.textContent=`✅ ${fileInput.files[0].name}`}}document.getElementById('analysisForm').addEventListener('submit',async(e)=>{e.preventDefault();const dateStart=document.getElementById('dateStart').value;const dateEnd=document.getElementById('dateEnd').value;const competitors=Array.from(document.getElementById('competitors').selectedOptions).map(o=>o.value);const file=document.getElementById('mordorFile').files[0];if(!dateStart||!dateEnd){showError('Please select both dates');return}if(!competitors.length){showError('Please select at least one competitor');return}if(!file){showError('Please upload an Excel file');return}document.getElementById('progressSection').classList.add('active');document.getElementById('submitBtn').disabled=true;const formData=new FormData();formData.append('dateStart',dateStart);formData.append('dateEnd',dateEnd);competitors.forEach(c=>formData.append('competitors',c));formData.append('mordor_file',file);try{const response=await fetch('/api/analyze',{method:'POST',body:formData});if(!response.ok){const error=await response.json();throw new Error(error.error||'Analysis failed')}const blob=await response.blob();const url=URL.createObjectURL(blob);document.getElementById('downloadLink').href=url;document.getElementById('progressSection').classList.remove('active');document.getElementById('resultsSection').classList.add('active');showSuccess('✅ Analysis complete! Download your report.')}catch(error){document.getElementById('progressSection').classList.remove('active');showError(`Error: ${error.message}`)}finally{document.getElementById('submitBtn').disabled=false}});function showError(msg){const el=document.getElementById('errorMessage');el.textContent=msg;el.style.display='block'}function showSuccess(msg){const el=document.getElementById('successMessage');el.textContent=msg;el.style.display='block'}</script></body></html>'''

@app.errorhandler(500)
def error_handler(e):
    return jsonify({'error': 'Server error'}), 500

if __name__ == '__main__':
    app.run(debug=False, host='0.0.0.0', port=int(os.environ.get('PORT', 3000)))
